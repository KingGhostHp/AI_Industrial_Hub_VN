#!/usr/bin/env python3
"""Build a consolidated GeoJSON of all Vietnamese industrial zones/clusters."""

from __future__ import annotations

import argparse
import html
import json
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import requests
from dotenv import load_dotenv

# Load environment variables from .env
load_dotenv()

DEFAULT_TOKEN = os.getenv('MAPBOX_ACCESS_TOKEN', '')

GEOCODING_URL = "https://api.mapbox.com/geocoding/v5/mapbox.places/{query}.json"
STOP_LABELS = [
    "Thời gian",
    "Thoi gian",
    "Tổng diện tích",
    "Tong dien tich",
    "Giá thuê",
    "Giá",
    "Gia",
    "Mật độ",
    "Mat do",
    "Hệ thống",
    "He thong",
    "Ngành nghề",
    "Nganh nghe",
    "Ưu đãi",
    "Uu dai",
    "Liên hệ",
    "Lien he",
    "Thu gọn",
    "TÌM KIẾM",
    "Tim kiem",
    "Pháp lý",
    "Phap ly",
    "Mã số",
    "Ma so",
]

MAX_QUERY_LEN = 240

ROOT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT_DIR / "data"
TRACKASIA_EXCEL_CHOICES = (
    "IIPVietNam_final.xlsx",
    "IIPVietNam_trackasia_fixed.xlsx",
)
GENERIC_NAME_KEYS = {
    "khu cong nghiep",
    "cum cong nghiep",
    "cụm công nghiệp",
    "khu công nghiệp",
    "khu cong nghiep/ccn",
    "cum cong nghiep/kcn",
}
MANUAL_GEOCODE_QUERIES = {
    "https://iipvietnam.com/khu-cong-nghiep-dat-cuoc-binh-duong-957388.html": "Khu công nghiệp Đất Cuốc, Bắc Tân Uyên, Bình Dương, Vietnam",
    "https://iipvietnam.com/khu-cong-nghiep-dat-do-1-ba-ria-vung-tau-844134.html": "Khu công nghiệp Đất Đỏ 1, Đất Đỏ, Bà Rịa - Vũng Tàu, Vietnam",
    "https://iipvietnam.com/khu-cong-nghiep-hung-phu-thai-binh-835537.html": "Cụm công nghiệp Hưng Phú, Hưng Hà, Thái Bình, Vietnam",
    "https://iipvietnam.com/khu-cong-nghiep-huu-thanh-long-an-795478.html": "Khu công nghiệp Hựu Thạnh, Đức Hòa, Long An, Vietnam",
}
PROVINCE_LOOKUP = None


def clean_text(value: Optional[str]) -> str:
    if not value:
        return ""
    text = html.unescape(str(value))
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,;.-")


def extract_field(blob: str, label: str) -> str:
    blob = blob or ""
    lower_blob = blob.lower()
    label_lower = label.lower()
    idx = lower_blob.find(label_lower)
    if idx == -1:
        return ""
    start = idx + len(label)
    remainder = blob[start:]
    remainder_lower = lower_blob[start:]
    stop_pos = len(remainder)
    for stop in STOP_LABELS:
        stop_lower = stop.lower()
        pos = remainder_lower.find(stop_lower)
        if pos != -1 and pos < stop_pos:
            stop_pos = pos
    return clean_text(remainder[:stop_pos])


def clamp_tokens(text: str, max_tokens: int = 20) -> str:
    tokens = text.split()
    if len(tokens) <= max_tokens:
        return text
    return " ".join(tokens[:max_tokens])


def resolve_data_file(filename: str) -> Path:
    """Locate supporting data files even when the repo is nested elsewhere."""
    candidates = [
        DATA_DIR / filename,
        ROOT_DIR.parent / "data" / filename,
        ROOT_DIR.parent / "IIPMAP" / "data" / filename,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Không tìm thấy file dữ liệu '{filename}' trong {candidates}")


def load_trackasia_coords(path: Optional[Path]) -> Dict[str, Dict[str, float]]:
    if not path or not path.exists():
        return {}
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        print("⚠️  Thiếu thư viện openpyxl nên bỏ qua Excel tọa độ có sẵn.", file=sys.stderr)
        return {}
    coords: Dict[str, Dict[str, float]] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header = [value.strip() if isinstance(value, str) else value for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    try:
        url_idx = header.index("URL")
        lat_idx = header.index("Latitude")
        lng_idx = header.index("Longitude")
    except ValueError:
        print(f"⚠️  Không tìm thấy cột URL/Latitude/Longitude trong {path}", file=sys.stderr)
        wb.close()
        return {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        url = (row[url_idx] or "").strip()
        lat = row[lat_idx]
        lng = row[lng_idx]
        if not url or lat is None or lng is None:
            continue
        coords[url] = {
            "lat": float(lat),
            "lng": float(lng),
            "source": path.name,
        }
    wb.close()
    return coords


def normalize_text(value: str) -> str:
    if not value:
        return ""
    value = unicodedata.normalize("NFD", value)
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    value = value.lower()
    value = re.sub(r"[^a-z0-9\s-]", " ", value)
    value = value.replace("-", " ")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def slug_from_url(url: str) -> str:
    if not url:
        return ""
    slug = url.rstrip("/").split("/")[-1]
    if slug.endswith(".html"):
        slug = slug[:-5]
    slug = re.sub(r"-\d+$", "", slug)
    return slug


def slug_to_title(url: str) -> str:
    slug = slug_from_url(url)
    if not slug:
        return ""
    tokens = slug.replace("-", " ").split()
    words = []
    for token in tokens:
        if token.isdigit():
            continue
        if token.lower() in {"kcn", "ccn"}:
            words.append(token.upper())
        elif re.match(r"^[0-9]+[a-z]+$", token.lower()):
            words.append(token.upper())
        else:
            words.append(token.capitalize())
    return " ".join(words)


def load_province_lookup() -> Dict[str, str]:
    global PROVINCE_LOOKUP
    if PROVINCE_LOOKUP is not None:
        return PROVINCE_LOOKUP
    lookup: Dict[str, str] = {}
    geo_path = Path(__file__).parent / "data" / "vn_provinces_34.geojson"
    if geo_path.exists():
        try:
            data = json.loads(geo_path.read_text(encoding="utf-8"))
            for feature in data.get("features", []):
                name = (feature.get("properties") or {}).get("NAME_1")
                if not name:
                    continue
                pretty_name = format_province_label(name)
                normalized = normalize_text(
                    name.replace("Tỉnh", "").replace("Thành phố", "")
                )
                if normalized:
                    lookup[normalized] = pretty_name
                    condensed = normalized.replace(" ", "")
                    lookup[condensed] = pretty_name
        except json.JSONDecodeError:
            pass
    manual_aliases = {
        "ho chi minh": "TP. Hồ Chí Minh",
        "ho chi minh city": "TP. Hồ Chí Minh",
        "tp ho chi minh": "TP. Hồ Chí Minh",
        "tphcm": "TP. Hồ Chí Minh",
        "ba ria vung tau": "Bà Rịa - Vũng Tàu",
        "thua thien hue": "Thừa Thiên - Huế",
    }
    lookup.update(manual_aliases)
    PROVINCE_LOOKUP = lookup
    return lookup


def format_province_label(raw: str) -> str:
    text = raw
    if " " not in text:
        text = re.sub(r"(?<!^)([A-ZĐ])", r" \1", text)
    text = text.replace("-", " - ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def guess_province_from_slug(url: str) -> str:
    slug = slug_from_url(url)
    normalized = normalize_text(slug)
    if not normalized:
        return ""
    tokens = normalized.split()
    lookup = load_province_lookup()
    for size in range(4, 0, -1):
        if len(tokens) < size:
            continue
        candidate = " ".join(tokens[-size:])
        options = {candidate, candidate.replace(" ", "")}
        for key in options:
            if key in lookup:
                return lookup[key]
    return ""


def enrich_feature_metadata(feature: Dict[str, str]) -> None:
    url = feature.get("url") or ""
    name = feature.get("name") or ""
    normalized_name = normalize_text(name)
    if not name or normalized_name in GENERIC_NAME_KEYS:
        slug_name = slug_to_title(url)
        if slug_name:
            feature["name"] = slug_name
    province = feature.get("province") or ""
    if not province or "không rõ" in province.lower():
        guessed = guess_province_from_slug(url)
        if guessed:
            feature["province"] = guessed
    if not feature.get("address") and feature.get("province"):
        feature["address"] = feature["province"]


def attach_precomputed_coords(features: Iterable[Dict[str, str]], coords: Dict[str, Dict[str, float]]) -> int:
    matched = 0
    if not coords:
        return matched
    for feature in features:
        url = (feature.get("url") or "").strip()
        record = coords.get(url)
        if not record:
            continue
        feature["lat"] = record["lat"]
        feature["lng"] = record["lng"]
        feature["coordinate_source"] = record.get("source", "precomputed")
        matched += 1
    return matched


def detect_excel_path(explicit: Optional[Path]) -> Optional[Path]:
    if explicit:
        return explicit if explicit.exists() else None
    candidates = []
    for name in TRACKASIA_EXCEL_CHOICES:
        candidates.extend(
            [
                ROOT_DIR / "data" / name,
                ROOT_DIR / "data" / "export" / name,
                ROOT_DIR.parent / name,
                ROOT_DIR.parent / "data" / name,
                ROOT_DIR.parent / "IIPMAP" / "data" / "export" / name,
            ]
        )
    for candidate in candidates:
        if candidate and candidate.exists():
            return candidate
    return None


def load_iipmap() -> List[Dict[str, str]]:
    path = resolve_data_file("export_ready_ultimate.records.jsonl")
    features = []
    seen = set()
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            row = json.loads(line)
            url = row.get("url")
            if not url or url in seen:
                continue
            seen.add(url)
            features.append(
                {
                    "id": url,
                    "url": url,
                    "name": clean_text(row.get("tên") or row.get("name")),
                    "kind": clean_text(row.get("loại")),
                    "province": clean_text(row.get("tỉnh")),
                    "address": clean_text(row.get("địa chỉ")),
                    "source": "iipmap.com",
                }
            )
    return features


def load_iipvietnam() -> List[Dict[str, str]]:
    path = resolve_data_file("export_ready_63.json")
    text = path.read_text(encoding="utf-8-sig")
    payload = json.loads(text)
    features: List[Dict[str, str]] = []
    seen = set()
    for province in payload.get("provinces", []):
        province_name = clean_text(province.get("name"))
        for kind in ("KCN", "CCN"):
            for entry in province.get(kind, []) or []:
                url = entry.get("url")
                if not url or url in seen:
                    continue
                seen.add(url)
                content = entry.get("content", "")
                vietnam_name = extract_field(content, "Tên:")
                fallback_name = clean_text(entry.get("name"))
                final_name = vietnam_name or fallback_name
                address = extract_field(content, "Địa chỉ:") or extract_field(content, "Dia chi:")
                features.append(
                    {
                        "id": url,
                        "url": url,
                        "name": final_name,
                        "kind": kind,
                        "province": province_name,
                        "address": clean_text(address or fallback_name),
                        "source": "iipvietnam.com",
                    }
                )
    return features


def geocode(
    query: str,
    *,
    token: str,
    country: str,
    types: str,
    limit: int,
    session: requests.Session,
) -> Dict[str, float]:
    url = GEOCODING_URL.format(query=requests.utils.quote(query))
    params = {
        "access_token": token,
        "country": country,
        "types": types,
        "limit": limit,
    }
    resp = session.get(url, params=params, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    features = data.get("features") or []
    if not features:
        return {}
    top = features[0]
    context_country = next(
        (c for c in top.get("context", []) if c.get("id", "").startswith("country")),
        {},
    )
    if country and context_country.get("short_code") != country.lower():
        return {}
    return {
        "match_name": top.get("place_name"),
        "confidence": top.get("relevance"),
        "lng": top["center"][0],
        "lat": top["center"][1],
    }


def geocode_features(
    features: Iterable[Dict[str, str]],
    *,
    token: str,
    country: str,
    types: str,
    limit: int,
    delay: float,
    cache_path: Path,
) -> Dict[str, Dict[str, float]]:
    cache: Dict[str, Dict[str, float]] = {}
    if cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            pass

    session = requests.Session()
    session.headers.update({"User-Agent": "iipmap-geojson-builder/1.0"})

    updated = False
    dirty = 0
    results: Dict[str, Dict[str, float]] = {}

    for feature in features:
        fid = feature["id"]
        manual_query = MANUAL_GEOCODE_QUERIES.get(fid)
        if fid in cache and not manual_query:
            results[fid] = cache[fid]
            continue
        preset_lng = feature.get("lng")
        preset_lat = feature.get("lat")
        if isinstance(preset_lng, (int, float)) and isinstance(preset_lat, (int, float)):
            result = {
                "match_name": feature.get("coordinate_source") or "precomputed",
                "confidence": 1.0,
                "lng": preset_lng,
                "lat": preset_lat,
                "query": "precomputed",
            }
            cache[fid] = result
            results[fid] = result
            updated = True
            dirty += 1
            if dirty >= 25:
                cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
                dirty = 0
                updated = False
            continue

        name = feature.get("name") or ""
        address = feature.get("address") or ""
        if address:
            address_parts = [seg.strip() for seg in address.split(",") if seg.strip()]
            if len(address_parts) > 3:
                address = ", ".join(address_parts[-3:])
        province = feature.get("province") or ""
        parts = [name]
        if address and address.lower() not in name.lower():
            parts.append(address)
        elif province:
            parts.append(province)
        parts.append("Vietnam")
        query = ", ".join([p for p in parts if p])
        if len(query) > MAX_QUERY_LEN:
            query = query[:MAX_QUERY_LEN]
        query = clamp_tokens(query)
        if not query:
            continue

        attempt = 0
        result = {}
        while attempt < 3:
            attempt += 1
            try:
                current_query = manual_query or query
                result = geocode(
                    current_query,
                    token=token,
                    country=country,
                    types=types,
                    limit=limit,
                    session=session,
                )
                if result:
                    if manual_query:
                        result["query"] = manual_query
                    break
            except requests.RequestException as exc:
                print(f"⚠️  Geocoding failed for '{name}' (attempt {attempt}): {exc}", file=sys.stderr)
                time.sleep(min(5, delay * 5))
        if not result and province and address:
            alt_query = clamp_tokens(f"{name}, {province}, Vietnam")
            try:
                result = geocode(
                    alt_query,
                    token=token,
                    country=country,
                    types=types,
                    limit=limit,
                    session=session,
                )
            except requests.RequestException:
                result = {}
        if result:
            result["query"] = query
            cache[fid] = result
            results[fid] = result
            updated = True
            dirty += 1
            print(f"✅ {name} -> {result['lng']:.6f}, {result['lat']:.6f} ({result['confidence']})")
        else:
            cache[fid] = {}
            results[fid] = {}
            updated = True
            dirty += 1
            print(f"❌ No match for '{name}' ({feature['source']})", file=sys.stderr)
        time.sleep(delay)
        if dirty >= 25:
            cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
            dirty = 0
            updated = False

    if updated or dirty:
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


def build_geojson(
    features: List[Dict[str, str]],
    matches: Dict[str, Dict[str, float]],
) -> Dict[str, object]:
    geo_features = []
    for feature in features:
        match = matches.get(feature["id"]) or {}
        if not match or "lng" not in match or "lat" not in match:
            continue
        props = {
            "name": feature.get("name"),
            "kind": feature.get("kind"),
            "source": feature.get("source"),
            "province": feature.get("province"),
            "address": feature.get("address"),
            "url": feature.get("url"),
            "match_name": match.get("match_name"),
            "confidence": match.get("confidence"),
        }
        geo_features.append(
            {
                "type": "Feature",
                "geometry": {
                    "type": "Point",
                    "coordinates": [match["lng"], match["lat"]],
                },
                "properties": props,
            }
        )
    return {"type": "FeatureCollection", "features": geo_features}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--token", default=DEFAULT_TOKEN, help="Mapbox access token")
    parser.add_argument("--country", default="vn", help="ISO country filter (default: vn)")
    parser.add_argument("--types", default="poi,place", help="Comma list of Mapbox feature types")
    parser.add_argument("--limit", type=int, default=3, help="Number of candidates per query")
    parser.add_argument("--delay", type=float, default=0.3, help="Delay between geocoding requests")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "data" / "industrial_zones.geojson",
        help="Destination GeoJSON path",
    )
    parser.add_argument(
        "--cache",
        type=Path,
        default=Path(__file__).parent / "data" / "geocoding_cache.json",
        help="Cache file to avoid re-geocoding the same entries",
    )
    parser.add_argument(
        "--excel-coords",
        type=Path,
        help="Optional Excel file chứa sẵn cột Latitude/Longitude (mặc định: auto detect)",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    dataset = []
    per_source = Counter()
    for loader in (load_iipmap, load_iipvietnam):
        chunk = loader()
        dataset.extend(chunk)
        if chunk:
            per_source[chunk[0]["source"]] += len(chunk)
    print(f"Loaded {len(dataset)} industrial/cụm CN entries.")
    for source, count in per_source.items():
        print(f"  - {source}: {count}")

    for feature in dataset:
        enrich_feature_metadata(feature)

    excel_path = detect_excel_path(args.excel_coords)
    coords_by_url = load_trackasia_coords(excel_path)
    if coords_by_url:
        matched = attach_precomputed_coords(dataset, coords_by_url)
        print(f"Đã gắn tọa độ sẵn cho {matched} mục từ {excel_path}")
    else:
        print("Không tìm thấy file Excel chứa tọa độ sẵn. Toàn bộ điểm sẽ gọi Mapbox Geocoding.")

    matches = geocode_features(
        dataset,
        token=args.token,
        country=args.country,
        types=args.types,
        limit=args.limit,
        delay=args.delay,
        cache_path=args.cache,
    )

    geojson = build_geojson(dataset, matches)
    args.output.write_text(json.dumps(geojson, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved {len(geojson['features'])} features to {args.output}")


if __name__ == "__main__":
    main()
